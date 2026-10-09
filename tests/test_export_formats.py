import csv
import sys
import unittest
from io import BytesIO, StringIO
from pathlib import Path
from openpyxl import Workbook, load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend.features.pubmed_scraper.export_formats import convert_workbook

class ExportTests(unittest.TestCase):
    def setUp(self):
        workbook = Workbook()
        data = workbook.active; data.title = 'PubMed Data'
        data.append(['Author', 'ID', 'No. Publications', 'Title'])
        data.append(['José', '00123', 7, 'Study, with commas'])
        data.append(['=literal', None, 0, 'line\nbreak'])
        workbook.create_sheet('Names Mapping').append(['Alias', 'Name'])
        output = BytesIO(); workbook.save(output); self.raw = output.getvalue()

    def test_excel_all_sheets_and_text(self):
        workbook = load_workbook(convert_workbook(self.raw, 'xlsx'))
        self.assertEqual(workbook.sheetnames, ['PubMed Data', 'Names Mapping'])
        self.assertEqual(workbook['PubMed Data']['B2'].value, '00123')
        self.assertEqual(workbook['PubMed Data']['C2'].value, '7')
        self.assertEqual(workbook['PubMed Data']['A3'].data_type, 's')
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertEqual(cell.number_format, '@')
                    if cell.value is not None: self.assertEqual(cell.data_type, 's')

    def test_csv_unicode_quotes_newlines(self):
        raw = convert_workbook(self.raw, 'csv').getvalue()
        rows = list(csv.reader(StringIO(raw.decode('utf-8-sig'))))
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1], ['José', '00123', '7', 'Study, with commas'])
        self.assertEqual(rows[2][3], 'line\nbreak')

    def test_invalid_format(self):
        with self.assertRaises(ValueError): convert_workbook(self.raw, 'pdf')

if __name__ == '__main__': unittest.main()
