"""Review and apply PubMed potential-name mappings in an in-memory workbook."""

from __future__ import annotations

import json
import re
import unicodedata
from collections import OrderedDict
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook


def normalize_name(value: object) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(character for character in value if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def _sheet_name(workbook, requested: str) -> str | None:
    return next((name for name in workbook.sheetnames if name.lower() == requested.lower()), None)


def _name_columns(frame: pd.DataFrame) -> tuple[str, str] | None:
    candidates = [column for column in frame.columns if any(word in str(column).lower() for word in ("name", "author"))]
    if len(candidates) >= 2:
        return str(candidates[0]), str(candidates[1])
    if len(frame.columns) >= 2:
        return str(frame.columns[0]), str(frame.columns[1])
    return None


def _author_column(frame: pd.DataFrame) -> str | None:
    return next((str(column) for column in frame.columns if str(column).strip().lower() in {"author", "name"}), None)


def _evidence(frame: pd.DataFrame, author_column: str | None, name: str) -> dict:
    if not author_column or author_column not in frame.columns:
        return {"publications": 0, "affiliations": []}
    target = normalize_name(name)
    rows = frame[frame[author_column].map(normalize_name) == target]
    fields = [column for column in ("HCO", "Department", "City", "Affiliations") if column in rows.columns]
    affiliations = []
    for _, row in rows[fields].fillna("").iterrows():
        value = " · ".join(str(row[column]).strip() for column in fields if str(row[column]).strip())
        if value and value not in affiliations:
            affiliations.append(value)
    publications = rows["URL"].nunique() if "URL" in rows.columns else len(rows)
    title_column = next((column for column in rows.columns if str(column).strip().lower() in {"title", "article title", "publication title"}), None)
    titles = list(dict.fromkeys(str(value).strip() for value in rows[title_column] if str(value).strip())) if title_column else []
    return {"publications": int(publications), "affiliations": affiliations[:12], "titles": titles[:12]}


def potential_mapping_review(raw: bytes) -> dict:
    """Return grouped potential name pairs and evidence from the source workbook."""
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    potential_sheet = _sheet_name(workbook, "Potential Mappings")
    data_sheet = _sheet_name(workbook, "PubMed Data")
    if not potential_sheet or not data_sheet:
        return {"groups": [], "pairs": [], "message": "This workbook has no Potential Mappings sheet."}
    potential = pd.read_excel(BytesIO(raw), sheet_name=potential_sheet, dtype=str).fillna("")
    data = pd.read_excel(BytesIO(raw), sheet_name=data_sheet, dtype=str).fillna("")
    columns = _name_columns(potential)
    if not columns:
        return {"groups": [], "pairs": [], "message": "Potential Mappings does not contain two name columns."}
    left_column, right_column = columns
    author_column = _author_column(data)
    grouped: OrderedDict[str, dict] = OrderedDict()
    for _, row in potential.iterrows():
        left, right = str(row[left_column]).strip(), str(row[right_column]).strip()
        if not left or not right or normalize_name(left) == normalize_name(right):
            continue
        key = "|".join(sorted((normalize_name(left), normalize_name(right))))
        if key not in grouped:
            grouped[key] = {"id": str(len(grouped) + 1), "name_a": left, "name_b": right}
    pairs = []
    for pair in grouped.values():
        pairs.append({**pair, "evidence_a": _evidence(data, author_column, pair["name_a"]), "evidence_b": _evidence(data, author_column, pair["name_b"])})
    # Connected components place every overlapping suggestion in one review.
    names = OrderedDict()
    adjacency = {}
    for pair in pairs:
        a, b = normalize_name(pair["name_a"]), normalize_name(pair["name_b"])
        names.setdefault(a, pair["name_a"])
        names.setdefault(b, pair["name_b"])
        adjacency.setdefault(a, set()).add(b)
        adjacency.setdefault(b, set()).add(a)
    groups, visited = [], set()
    for root in names:
        if root in visited:
            continue
        stack, component = [root], set()
        while stack:
            key = stack.pop()
            if key in component:
                continue
            component.add(key)
            stack.extend(adjacency[key] - component)
        visited.update(component)
        members = [names[key] for key in names if key in component]
        groups.append({
            "id": str(len(groups) + 1),
            "names": [{"name": name, "evidence": _evidence(data, author_column, name)} for name in members],
            "suggested_name": max(members, key=len),
            "pairs": [pair for pair in pairs if normalize_name(pair["name_a"]) in component],
        })
    return {"groups": groups, "pairs": pairs, "message": ""}


def _clean_review_sheets(workbook) -> None:
    review_name = _sheet_name(workbook, "Name Review")
    if review_name:
        del workbook[review_name]
    potential_name = _sheet_name(workbook, "Potential Mappings")
    if potential_name:
        worksheet = workbook[potential_name]
        has_rows = any(
            any(value is not None and str(value).strip() for value in row)
            for row in worksheet.iter_rows(min_row=2, values_only=True)
        )
        if not has_rows:
            del workbook[potential_name]


def apply_name_decisions(raw: bytes, decisions_json: str) -> BytesIO:
    """Merge approved aliases and move them from Potential to Names Mapping."""
    try:
        decisions = json.loads(decisions_json)
    except json.JSONDecodeError as exc:
        raise ValueError("Name-review decisions are not valid JSON.") from exc
    if not isinstance(decisions, list):
        raise ValueError("Name-review decisions must be a list.")

    if any(isinstance(item, dict) and "group_id" in item for item in decisions) and not all(isinstance(item, dict) and "group_id" in item for item in decisions):
        raise ValueError("Do not mix grouped and pair review decisions.")

    if not decisions:
        workbook = load_workbook(BytesIO(raw))
        _clean_review_sheets(workbook)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    review = potential_mapping_review(raw)
    groups = {group["id"]: group for group in review["groups"]}
    replacements: dict[str, str] = {}
    merged_pairs: dict[str, tuple[str, str]] = {}
    resolved_pairs = set()
    reviewed_groups = set()
    for decision in decisions:
        if not isinstance(decision, dict):
            raise ValueError("Each name-review decision must be an object.")
        if "group_id" in decision:
            group_id = str(decision["group_id"])
            if group_id not in groups or group_id in reviewed_groups:
                raise ValueError("Unknown or repeated review group.")
            reviewed_groups.add(group_id)
            group = groups[group_id]
            allowed = {normalize_name(item["name"]): item["name"] for item in group["names"]}
            merges = decision.get("merges", [])
            if not isinstance(merges, list):
                raise ValueError("Group merges must be a list.")
            used = set()
            for merge in merges:
                if not isinstance(merge, dict) or not isinstance(merge.get("names"), list):
                    raise ValueError("Each merge must contain a list of names.")
                selected = [normalize_name(name) for name in merge["names"]]
                canonical = normalize_name(merge.get("canonical_name", ""))
                if len(selected) < 2 or len(set(selected)) != len(selected) or not set(selected) <= allowed.keys() or canonical not in selected:
                    raise ValueError("Select at least two names from the group and keep one selected name.")
                if used.intersection(selected):
                    raise ValueError("A name cannot belong to two merges.")
                used.update(selected)
                for key in selected:
                    replacements[key] = allowed[canonical]
            for pair in group["pairs"]:
                resolved_pairs.add("|".join(sorted((normalize_name(pair["name_a"]), normalize_name(pair["name_b"])))))
            for key, name in allowed.items():
                final = replacements.get(key, name)
                if normalize_name(final) != key:
                    merged_pairs[key] = (name, final)
        else:
            # Support older clients, but reject contradictory pair decisions.
            a, b = str(decision.get("name_a", "")).strip(), str(decision.get("name_b", "")).strip()
            key = "|".join(sorted((normalize_name(a), normalize_name(b))))
            valid = {"|".join(sorted((normalize_name(pair["name_a"]), normalize_name(pair["name_b"])))) for pair in review["pairs"]}
            if key not in valid or decision.get("action") not in {"merge", "separate"}:
                raise ValueError("Invalid pair decision.")
            resolved_pairs.add(key)
            if decision["action"] == "merge":
                canonical = str(decision.get("canonical_name", "")).strip()
                if normalize_name(canonical) not in {normalize_name(a), normalize_name(b)}:
                    raise ValueError("Keep one of the names being merged.")
                for name in (a, b):
                    normalized = normalize_name(name)
                    if normalized in replacements and normalize_name(replacements[normalized]) != normalize_name(canonical):
                        raise ValueError("Overlapping pair decisions conflict. Use grouped name review.")
                    replacements[normalized] = canonical
                other = b if normalize_name(canonical) == normalize_name(a) else a
                merged_pairs[normalize_name(other)] = (other, canonical)

    workbook = load_workbook(BytesIO(raw))
    data_sheet = _sheet_name(workbook, "PubMed Data")
    if not data_sheet:
        raise ValueError("The workbook has no PubMed Data sheet.")
    # Apply the same approved names to publication authors and collaborator rows.
    for sheet_name in (data_sheet, _sheet_name(workbook, "Collaborators")):
        if not sheet_name:
            continue
        worksheet = workbook[sheet_name]
        name_columns = [cell.column for cell in worksheet[1] if str(cell.value).strip().lower() in {"author", "name"}]
        for row in range(2, worksheet.max_row + 1):
            for column in name_columns:
                cell = worksheet.cell(row=row, column=column)
                replacement = replacements.get(normalize_name(cell.value))
                if replacement:
                    cell.value = replacement

        headers = {str(cell.value).strip().lower(): cell.column for cell in worksheet[1]}
        author_col = headers.get("author") or headers.get("name")
        url_col = headers.get("url")
        count_col = headers.get("no. publications")
        if replacements and author_col and url_col and count_col:
            urls = {}
            for row in range(2, worksheet.max_row + 1):
                name = normalize_name(worksheet.cell(row, author_col).value)
                url = worksheet.cell(row, url_col).value
                if url is not None and str(url).strip():
                    urls.setdefault(name, set()).add(str(url))
            for row in range(2, worksheet.max_row + 1):
                name = normalize_name(worksheet.cell(row, author_col).value)
                worksheet.cell(row, count_col).value = len(urls.get(name, set()))

    # Keep the original Names Mapping layout. The first two fields are the
    # alias and canonical name; the third field marks decisions made here.
    names_mapping_name = _sheet_name(workbook, "Names Mapping")
    if names_mapping_name:
        names_mapping_sheet = workbook[names_mapping_name]
        if names_mapping_sheet.max_row == 0:
            names_mapping_sheet.append(["Source name", "Mapped name", "Mapping type"])
    else:
        names_mapping_sheet = workbook.create_sheet("Names Mapping")
        names_mapping_sheet.append(["Source name", "Mapped name", "Mapping type"])
    for row in names_mapping_sheet.iter_rows(min_row=2):
        if len(row) >= 2:
            replacement = replacements.get(normalize_name(row[1].value))
            if replacement:
                row[1].value = replacement
    for source_name, canonical_name in merged_pairs.values():
        names_mapping_sheet.append([source_name, canonical_name, "Potential Map"])

    # Every suggestion in a confirmed group is resolved, including keep-separate decisions.
    potential_mapping_name = _sheet_name(workbook, "Potential Mappings")
    if potential_mapping_name and resolved_pairs:
        potential_sheet = workbook[potential_mapping_name]
        header_values = [cell.value for cell in potential_sheet[1]]
        potential_frame = pd.DataFrame(
            potential_sheet.iter_rows(min_row=2, values_only=True),
            columns=header_values,
        ).fillna("")
        pair_columns = _name_columns(potential_frame)
        if pair_columns:
            left_column, right_column = pair_columns
            keep_rows = []
            for _, row in potential_frame.iterrows():
                key = "|".join(sorted((normalize_name(row[left_column]), normalize_name(row[right_column]))))
                if key not in resolved_pairs:
                    keep_rows.append(row.tolist())
            potential_sheet.delete_rows(2, potential_sheet.max_row)
            for row in keep_rows:
                potential_sheet.append(row)

    _clean_review_sheets(workbook)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
