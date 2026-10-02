import csv
from io import BytesIO, StringIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _text(value):
    """Return an export-safe string without turning blanks into 'nan'."""
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value)


def _write_text_sheet(worksheet, headers, rows):
    """Write every cell as an explicit Excel text cell."""
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=_text(header))
        cell.font = Font(bold=True)
        cell.data_type = "s"

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=_text(value))
            cell.data_type = "s"

    for column_index, header in enumerate(headers, start=1):
        values = [_text(header)] + [_text(row[column_index - 1]) for row in rows[:500]]
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(len(value) for value in values) + 2,
            60,
        )


def append_unmapped_target_rows(
    rows,
    target_frame,
    target_columns,
    mapped_target_indices,
    source_column_count,
):
    mapped = {int(index) for index in mapped_target_indices}
    output_rows = list(rows)
    for target_index, target_row in target_frame.reset_index(drop=True).iterrows():
        if target_index in mapped:
            continue
        output_rows.append(
            [""] * source_column_count
            + [
                _text(value)
                for value in (target_row[column] for column in target_columns)
            ]
        )
    return output_rows


def create_excel_workbook(headers, rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Matches"
    _write_text_sheet(worksheet, headers, rows)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def create_text_excel_workbook(sheets):
    """Create a multi-sheet workbook where every value is stored as text."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, headers, rows in sheets:
        worksheet = workbook.create_sheet(title=sheet_name)
        _write_text_sheet(worksheet, list(headers), list(rows))

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def create_csv_file(headers, rows):
    text_output = StringIO(newline="")
    # A CSV format has no cell types. Quoting every string preserves the values
    # exactly in the file (including leading zeroes) for CSV-aware readers.
    writer = csv.writer(text_output, lineterminator="\r\n", quoting=csv.QUOTE_ALL)
    writer.writerow([_text(value) for value in headers])
    writer.writerows(
        [_text(value) for value in row]
        for row in rows
    )
    output = BytesIO(text_output.getvalue().encode("utf-8-sig"))
    output.seek(0)
    return output
