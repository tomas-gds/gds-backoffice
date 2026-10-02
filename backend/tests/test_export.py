import unittest
from io import BytesIO

from openpyxl import load_workbook

import pandas as pd

from backend.features.database_mapper.exporter import (
    append_unmapped_target_rows,
    create_csv_file,
    create_excel_workbook,
    create_text_excel_workbook,
)


class ExportTests(unittest.TestCase):
    def test_csv_export_preserves_unicode_and_quotes_delimiters(self):
        output = create_csv_file(
            ["Name", "Address"],
            [["Hôpital de la Charité", "Street, 1"], ["Gréoux-les-Bains", 'A "quoted" value']],
        )
        text = output.read().decode("utf-8-sig")
        self.assertIn("Hôpital de la Charité", text)
        self.assertIn('"Street, 1"', text)
        self.assertIn('"A ""quoted"" value"', text)

    def test_unused_file_b_rows_are_appended_with_blank_file_a_columns(self):
        rows = [["Source A", "Matched B", "Paris"]]
        target_frame = pd.DataFrame(
            [
                {"Name": "Matched B", "City": "Paris"},
                {"Name": "Hôpital de la Charité", "City": "Gréoux-les-Bains"},
            ]
        )

        result = append_unmapped_target_rows(
            rows,
            target_frame,
            ["Name", "City"],
            [0],
            1,
        )

        self.assertEqual(
            result,
            [
                ["Source A", "Matched B", "Paris"],
                ["", "Hôpital de la Charité", "Gréoux-les-Bains"],
            ],
        )

    def test_excel_export_preserves_unicode_and_text_values(self):
        output = create_excel_workbook(
            ["HCO Name", "City", "Identifier"],
            [
                ["Hôpital de la Charité", "Gréoux-les-Bains", "9.39294E+17"],
                ["A.I.M.T.", "Bouches-du-Rhône", "=literal value"],
            ],
        )

        workbook = load_workbook(BytesIO(output.read()), data_only=False)
        worksheet = workbook["Matches"]
        self.assertEqual(worksheet["A2"].value, "Hôpital de la Charité")
        self.assertEqual(worksheet["B2"].value, "Gréoux-les-Bains")
        self.assertEqual(worksheet["C2"].value, "9.39294E+17")
        self.assertEqual(worksheet["C3"].value, "=literal value")
        self.assertEqual(worksheet["C3"].data_type, "s")

    def test_duplicate_workbook_stores_leading_zeroes_as_text(self):
        output = create_text_excel_workbook(
            [
                ("File A", ["Postcode", "Identifier"], [["00123", "000456"]]),
                ("File B", ["Postcode"], [["04567"]]),
            ]
        )

        workbook = load_workbook(BytesIO(output.read()), data_only=False)
        self.assertEqual(workbook["File A"]["A2"].value, "00123")
        self.assertEqual(workbook["File A"]["A2"].data_type, "s")
        self.assertEqual(workbook["File B"]["A2"].value, "04567")
        self.assertEqual(workbook["File B"]["A2"].data_type, "s")


if __name__ == "__main__":
    unittest.main()
